from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE_TYPE

OUT = '/home/workspace/slide-decks/order-of-demolay/order-of-demolay-introduction.pptx'
W, H = 13.333, 7.5
navy = RGBColor(27, 54, 95)
blue = RGBColor(30, 68, 136)
gold = RGBColor(255, 207, 0)
cream = RGBColor(248, 247, 250)
muted = RGBColor(224, 229, 239)
white = RGBColor(255, 255, 255)
teal = RGBColor(0, 118, 175)
ink = RGBColor(15, 27, 45)
coral = RGBColor(172, 36, 42)
sky = RGBColor(114, 199, 246)
purple = RGBColor(102, 29, 52)
EMBLEM = '/home/workspace/slide-decks/order-of-demolay/assets/demolay-emblem-color.png'

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)

def tx(slide, text, x, y, w, h, size=20, color=white, bold=False, font='Lato', align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear(); tf.word_wrap = True; tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text; r.font.name = font; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return box

def img(slide, path, x, y, w=None, h=None, transparency=0):
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w) if w else None, height=Inches(h) if h else None)
    if transparency:
        pic.fill.transparency = transparency
    return pic

def bg(slide, color=navy):
    fill = slide.background.fill; fill.solid(); fill.fore_color.rgb = color
    panel = slide.shapes.add_shape(MSO_SHAPE.PARALLELOGRAM, Inches(9.7), Inches(-.8), Inches(4.4), Inches(9.2))
    panel.fill.solid(); panel.fill.fore_color.rgb = purple; panel.fill.transparency = 12; panel.line.fill.background()
    panel2 = slide.shapes.add_shape(MSO_SHAPE.PARALLELOGRAM, Inches(-1.2), Inches(6.65), Inches(4.3), Inches(1.3))
    panel2.fill.solid(); panel2.fill.fore_color.rgb = RGBColor(22, 53, 81); panel2.fill.transparency = 20; panel2.line.fill.background()
    for x, y, d, c in [(11.8, -.35, 1.8, blue), (-.5, 6.55, 1.4, gold), (10.85, 5.72, .42, gold), (11.48, 5.25, .22, teal), (12.05, 4.78, .14, blue)]:
        sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
        sh.fill.solid(); sh.fill.fore_color.rgb = c; sh.fill.transparency = 18; sh.line.fill.background()
    for i in range(4):
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(10.9+i*.35), Inches(5.85-i*.32), Inches(12.1+i*.18), Inches(4.9-i*.28))
        line.line.color.rgb = gold; line.line.transparency = 55; line.line.width = Pt(1.2)
    for i in range(9):
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(.55+i*.42), Inches(1.55), Inches(.55+i*.42), Inches(6.55))
        line.line.color.rgb = RGBColor(111, 145, 194); line.line.transparency = 80; line.line.width = Pt(.55)
    for x, y in [(1.05, 1.55), (2.8, 5.8), (4.9, 1.3), (7.1, 6.25), (8.15, 1.5)]:
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(.07), Inches(.07))
        dot.fill.solid(); dot.fill.fore_color.rgb = gold; dot.line.fill.background()

def footer(slide, n):
    tx(slide, f'{n:02d}  |  THE ORDER OF DeMOLAY', .65, 7.1, 5, .2, 9, muted, True)

def title(slide, kicker, head, n):
    tx(slide, kicker.upper(), .7, .48, 5, .25, 11, gold, True)
    tx(slide, head, .7, .82, 11.5, .9, 35, white, True, 'Bebas Neue')
    footer(slide, n)

def card(slide, x, y, w, h, heading, body, accent=blue, icon=None):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor(21, 43, 76); sh.line.color.rgb = RGBColor(65, 99, 135); sh.line.width = Pt(1.1)
    stripe=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(.08), Inches(h))
    stripe.fill.solid(); stripe.fill.fore_color.rgb = accent; stripe.line.fill.background()
    if icon:
        badge=slide.shapes.add_shape(MSO_SHAPE.HEXAGON, Inches(x+.2), Inches(y+.18), Inches(.48), Inches(.48))
        badge.fill.solid(); badge.fill.fore_color.rgb = accent; badge.line.fill.background()
        tx(slide, icon, x+.2, y+.25, .48, .25, 8 if len(icon) > 2 else 12, navy, True, align=PP_ALIGN.CENTER)
    tx(slide, heading, x+.25+(0.5 if icon else 0), y+.2, w-.5, .4, 18, white, True)
    tx(slide, body, x+.25, y+.78, w-.5, h-.95, 16, white, False)

def bullets(slide, items, x=.9, y=2.1, w=6.2, size=21):
    for i, item in enumerate(items):
        yy=y+i*.75
        sh=slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(yy+.1), Inches(.16), Inches(.16))
        sh.fill.solid(); sh.fill.fore_color.rgb = gold; sh.line.fill.background()
        tx(slide, item, x+.35, yy, w, .45, size, cream)

def add_slide(kicker, head, n):
    s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s); title(s,kicker,head,n); return s

def shield(slide, x, y, w, h, fill=blue, label=''):
    sh=slide.shapes.add_shape(MSO_SHAPE.HEXAGON, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=fill; sh.line.color.rgb=gold; sh.line.width=Pt(2)
    inner=slide.shapes.add_shape(MSO_SHAPE.HEXAGON, Inches(x+.13), Inches(y+.13), Inches(w-.26), Inches(h-.26))
    inner.fill.background(); inner.line.color.rgb=RGBColor(255,255,255); inner.line.transparency=65; inner.line.width=Pt(1)
    if label: tx(slide,label,x+.25,y+h*.38,w-.5,.5,18,white,True,align=PP_ALIGN.CENTER)
    return sh

# 1
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
tx(s,'THE ORDER OF',.78,1.15,6,.35,18,gold,True,'Bebas Neue'); tx(s,'DeMOLAY',.72,1.48,7,1.0,60,white,True,'Bebas Neue')
tx(s,'A place to lead, serve, belong, and become.',.78,2.75,8,.5,27,white)
tx(s,'An introduction for young men and the people who support them',.8,5.9,8,.35,16,muted)
img(s, EMBLEM, 9.45, 1.05, 2.8, 2.8)
tx(s,'THE ROAD\nBEGINS HERE',9.8,4.05,2.1,.7,15,gold,True,align=PP_ALIGN.CENTER); footer(s,1)

# 2
s=add_slide('Start here','What is DeMolay?',2); bullets(s,['A youth leadership organization','Built around friendship and service','Run by young people with adult support','A place to practice real responsibility'],.9,2.0,6.3,23)
card(s,8.0,2.0,4.3,3.65,'The short version','DeMolay helps young men develop character, confidence, and leadership through shared experiences—not lectures.',gold,'01')

# 3
s=add_slide('The experience','You do more than attend meetings',3)
for x1,y1,x2,y2 in [(2.15,3.7,4.35,3.7),(5.15,3.7,7.35,3.7),(8.15,3.7,10.35,3.7)]:
    connector=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)); connector.line.color.rgb=gold; connector.line.transparency=18; connector.line.width=Pt(2)
card(s,.8,2.0,2.75,3.6,'Lead','Plan events. Speak up. Take a role.',blue,'◆')
card(s,3.8,2.0,2.75,3.6,'Serve','Help your community. Make an impact.',teal,'♥')
card(s,6.8,2.0,2.75,3.6,'Belong','Build friendships across chapters.',gold,'●')
card(s,9.8,2.0,2.75,3.6,'Grow','Learn skills you can use anywhere.',RGBColor(213,111,88),'↗')

# 4
s=add_slide('The promise','Seven virtues. One direction.',4)
tx(s,'The Crown of Youth represents the ideals DeMolays are challenged to practice.',.8,1.8,8,.45,20,white)
img(s, EMBLEM, 10.85, 4.95, 1.45, 1.45)
virt=['Filial Love','Reverence','Courtesy','Comradeship','Fidelity','Cleanness','Patriotism']
for i,v in enumerate(virt):
    x=.95+(i%4)*3.05; y=2.65+(i//4)*1.3
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(.55),Inches(.55)); sh.fill.solid(); sh.fill.fore_color.rgb=gold; sh.line.fill.background()
    tx(s,str(i+1),x,y+.08,.55,.3,16,navy,True,align=PP_ALIGN.CENTER); tx(s,v,x+.72,y+.1,2.1,.3,16,white,True)
tx(s,'Not perfection. Practice.',.8,5.9,5,.4,25,gold,True)

# 5
s=add_slide('Where you fit','Your chapter is youth-led',5)
tx(s,'Young members do the leading.',.9,1.9,5.2,.55,27,white,True)
bullets(s,['Elect and serve as officers','Create programs and activities','Work together on decisions','Welcome and mentor new members'],.95,2.8,5.9,20)
card(s,7.4,1.95,4.75,3.8,'Adults make the experience safe','Advisors provide guidance, continuity, training, and support so young leaders can lead well.',teal,'+')

# 6
s=add_slide('Make it real','What might you actually do?',6)
items=[('Service projects','Turn concern into action.'),('Leadership events','Practice speaking, planning, and teamwork.'),('Social activities','Have fun with people who get it.'),('Ceremonies and traditions','Connect today’s choices to lasting ideals.')]
for i,(h,b) in enumerate(items): card(s,.9+(i%2)*6.1,2.0+(i//2)*2.0,5.5,1.45,h,b,[blue,teal,gold,RGBColor(213,111,88)][i],str(i+1))

# 7
s=add_slide('The payoff','Skills that travel with you',7)
tx(s,'DeMolay is practice for the rest of life.',.9,1.75,7,.45,24,white,True)
tx(s,'01',.95,5.65,1.1,.5,30,blue,True); tx(s,'02',4.0,5.65,1.1,.5,30,teal,True); tx(s,'03',7.05,5.65,1.1,.5,30,gold,True); tx(s,'04',10.1,5.65,1.1,.5,30,coral,True)
skills=[('Confidence','Speak clearly—even when nervous.'),('Leadership','Move a team from idea to action.'),('Character','Choose what is right when it is hard.'),('Connection','Build relationships that last.')]
for i,(h,b) in enumerate(skills):
    x=.95+i*3.05; sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(2.75),Inches(2.35),Inches(2.25)); sh.fill.solid(); sh.fill.fore_color.rgb=[RGBColor(18,52,92),RGBColor(18,62,69),RGBColor(78,60,25),RGBColor(73,42,41)][i]; sh.line.color.rgb=[blue,teal,gold,coral][i]; sh.line.width=Pt(1.5)
    tx(s,h,x+.2,3.05,1.95,.4,20,white,True,'Crimson Pro'); tx(s,b,x+.2,3.65,1.95,.8,16,white,False,'Lato')

# 8
s=add_slide('A practical question','Who can join?',8)
card(s,.9,2.0,3.55,3.3,'Typical age','Young men who have passed their 12th birthday and are not yet 21.',gold,'12→21')
card(s,4.9,2.0,3.55,3.3,'What matters','Character, moral responsibility, belief in God, and respect for others.',teal,'✓')
card(s,8.9,2.0,3.55,3.3,'Local details','Eligibility, sponsorship, fees, and activities can vary by jurisdiction.',blue,'?')
tx(s,'Ask your local chapter for the current details.',.9,5.95,7,.35,20,gold,True)

# 9
s=add_slide('The invitation','You do not have to arrive ready',9)
tx(s,'You only need enough curiosity to take the next step.',.9,1.8,10,.55,29,white,True)
steps=[('1','Visit','Meet the chapter and see what it feels like.'),('2','Try','Join an activity or service project.'),('3','Ask','Talk with a member, advisor, or parent.'),('4','Decide','Choose whether DeMolay fits you.')]
for i,(num,h,b) in enumerate(steps):
    x=.95+i*3.05; tx(s,num,x,3.05,.5,.55,34,gold,True); tx(s,h,x+.65,3.05,1.8,.4,20,white,True); tx(s,b,x+.65,3.65,1.95,.8,16,white)

# 10
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
tx(s,'YOUR NEXT MOVE',.8,1.0,5,.3,12,gold,True,'Bebas Neue'); tx(s,'Come see what\nyou can become.',.75,1.55,8,1.4,48,white,True,'Bebas Neue')
tx(s,'Talk with a local DeMolay chapter.\nBring your questions. Bring a friend.',.8,3.55,6,.8,25,white)
sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(8.3),Inches(2.0),Inches(3.6),Inches(2.2)); sh.fill.solid(); sh.fill.fore_color.rgb=gold; sh.line.fill.background(); tx(s,'START\nHERE',8.65,2.45,2.9,1.1,32,navy,True,align=PP_ALIGN.CENTER); tx(s,'demolay.org',8.9,4.65,2.4,.3,17,gold,True,align=PP_ALIGN.CENTER); footer(s,10)

prs.save(OUT)
print(OUT)
